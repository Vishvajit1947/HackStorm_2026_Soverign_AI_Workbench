"""doc.write produces documents a human can actually open (SS 26117).

The brief asks for approval notes and calculation sheets - "real deliverables ...
not just chat replies". This suite opens every file it writes with the library
that reads that format, because a file with the right extension and the right
magic bytes can still be one a reader cannot use: a spreadsheet whose numbers are
stored as text cannot be summed, and a Word file with no heading styles is a wall
of prose rather than a note somebody signs.
"""

from __future__ import annotations

import pytest

from kavach.tools.impl.deliverables import DeliverableError, doc_write
from kavach.tools.workspace import Workspace


@pytest.fixture()
def ws(tmp_path):
    return Workspace(tmp_path / "run_deliverable")


def test_a_word_file_opens_and_keeps_its_structure(ws):
    from docx import Document

    doc_write(ws, "note.docx", title="Approval Note",
              body="Background:\n\nThickness is 8.89 mm.\n\nRecommendation:\n\nReplace it.")
    d = Document(ws.root / "note.docx")
    styles = [(p.style.name, p.text) for p in d.paragraphs if p.text.strip()]

    # Word's OWN Title and Heading styles, not hand-formatted runs. The first
    # thing anyone does with an approval note is restyle it, and a document
    # whose headings are 16pt bold runs cannot be restyled from Word's gallery
    # at all. The title was Heading 1 before this module became a typesetter.
    assert ("Title", "Approval Note") in styles, styles
    assert any(s.startswith("Heading") and t == "Background" for s, t in styles), styles
    assert any(s.startswith("Heading") and t == "Recommendation" for s, t in styles), styles
    assert any(t == "Thickness is 8.89 mm." for _s, t in styles)


def test_a_spreadsheet_stores_numbers_as_numbers(ws):
    """A column of strings cannot be summed, which makes it not a calculation
    sheet however right it looks."""
    from openpyxl import load_workbook

    doc_write(ws, "calc.xlsx", title="NDT",
              body="Tag,Thickness_mm\n6-P-1203,8.89\n6-P-1204,9.12\n")
    sheet = load_workbook(ws.root / "calc.xlsx").active
    assert sheet.title == "NDT"
    values = list(sheet.iter_rows(values_only=True))
    # The header stays on ROW 1. Every consumer of a spreadsheet - pandas, a
    # macro, the next step of a plan - looks for the columns there, so the
    # document title goes to the sheet tab and the PRINT header rather than
    # taking a data row for a banner.
    assert sheet.freeze_panes == "A2", "the header scrolls out of sight"

    # A machine key is not a column heading anyone circulates. `Thickness_mm`
    # is written out with its unit; `Tag` was already readable and is untouched.
    assert values[0] == ("Tag", "Thickness (mm)"), values[0]

    assert values[1][1] == pytest.approx(8.89)
    assert isinstance(values[1][1], float), "a number was stored as text"
    assert values[1][0] == "6-P-1203", "a tag number was mangled into a number"


def test_a_deck_has_one_slide_per_block(ws):
    from pptx import Presentation

    doc_write(ws, "deck.pptx", title="Review",
              body="Findings\n- Below t-min\n\nActions\n- Replace impeller")
    prs = Presentation(ws.root / "deck.pptx")
    titles = [s.shapes.title.text for s in prs.slides if s.shapes.title]
    assert "Review" in titles and "Findings" in titles and "Actions" in titles


def test_the_extension_chooses_the_format(ws):
    assert doc_write(ws, "a.docx", body="x")["format"] == "docx"
    assert doc_write(ws, "b.xlsx", body="x")["format"] == "xlsx"
    assert doc_write(ws, "c.pptx", body="x")["format"] == "pptx"


def test_an_unsupported_extension_is_refused_with_the_alternative(ws):
    """Never a silently-wrong file: say what this tool does and name the tool
    that does the other thing."""
    with pytest.raises(DeliverableError, match="fs.write"):
        doc_write(ws, "notes.txt", body="x")


def test_a_deliverable_cannot_escape_the_workspace(ws):
    """doc.write goes through the same jail as fs.write - a new tool must not be
    a new way out."""
    from kavach.tools.workspace import WorkspaceViolation

    with pytest.raises((WorkspaceViolation, ValueError)):
        doc_write(ws, "../escape.docx", body="x")


def test_the_written_path_is_reported_relative_to_the_workspace(ws):
    result = doc_write(ws, "reports/note.docx", title="T", body="body")
    assert result["path"].replace("\\", "/") == "reports/note.docx"
    assert result["size_bytes"] > 0
    assert (ws.root / "reports" / "note.docx").exists()
