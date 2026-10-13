"""A file that opens is not the same thing as a deliverable.

From a report that generated .pptx files had "many gaps and sometimes less
data, single 1 two 2 words data, not exact findings". The decks opened, had the
right title and the right section headings, and said nothing - because the body
they were given was a template whose `{{stepN}}` reference resolved to almost
nothing, and a heading with no content under it still produced a slide.

Nothing here refuses to write the file: a thin deliverable is still the
deliverable, and refusing would lose the little there is. What changes is that
the thinness is part of the RESULT - in `warnings`, in the "Rendering notes"
inside the document, and therefore in the panel.
"""

from __future__ import annotations

import pytest

from kavach.tools.impl.deliverables import doc_write
from kavach.tools.workspace import Workspace

REAL_FINDING = (
    "The pump P-101A shows bearing vibration of 4.8 mm/s RMS against an "
    "ISO 10816-3 Zone B limit of 4.5 mm/s, measured on 02-MAR-2026 at the "
    "drive end in the horizontal direction."
)


@pytest.fixture()
def ws(tmp_path):
    return Workspace(tmp_path / "run")


# ------------------------------------------------------------- blank slides

def test_a_heading_with_no_content_does_not_become_a_blank_slide(ws):
    """The guard used to be "no heading AND no content", so a section with a
    title and nothing under it produced a slide with a title and an empty box."""
    out = doc_write(ws, "deck.pptx", title="Review",
                    body="Findings:\n\nRecommendation:\n\nActions:")
    assert out["empty_sections"] == ["Findings", "Recommendation", "Actions"]
    assert out["warnings"], "three empty sections and nothing said about it"


def test_an_empty_section_says_so_on_the_slide_itself(ws):
    """Not dropped, either. A section the model meant to fill and did not is a
    fact about the run, and it belongs in the deck a person opens."""
    from pptx import Presentation

    doc_write(ws, "deck.pptx", title="Review", body="Findings:\n\nActions:")
    text = "\n".join(
        shape.text
        for slide in Presentation(str(ws.resolve("deck.pptx"))).slides
        for shape in slide.shapes
        if getattr(shape, "text", "")
    )
    assert "no content was written for this section" in text
    assert "Rendering notes" in text


def test_a_full_deck_carries_no_empty_sections(ws):
    out = doc_write(
        ws, "good.pptx", title="Review",
        body=f"Findings:\n\n{REAL_FINDING}\n\nActions:\n\n- Re-measure after realignment\n- Inspect the mechanical seal",
    )
    assert out["empty_sections"] == []
    assert out["warnings"] == []


# -------------------------------------------------------- an outline, not a doc

def test_an_outline_is_reported_as_an_outline(ws):
    out = doc_write(ws, "hollow.docx", title="Note",
                    body="# Findings\n\n# Recommendation\n\n# Actions")
    assert out["content_chars"] == 0
    assert any("outline" in w for w in out["warnings"])


def test_two_words_per_section_is_still_an_outline(ws):
    """The reported case exactly: sections that are filled, technically."""
    out = doc_write(ws, "thin.pptx", title="Review",
                    body="Findings:\n\nPump ok\n\nActions:\n\nCheck seal")
    assert any("outline" in w for w in out["warnings"])


def test_a_short_but_real_note_is_not_flagged(ws):
    """"How short is too short" has no absolute answer - a nine-character note
    under one heading is a legitimate document. What is wrong in the hollow case
    is the RATIO: sections laid out and not filled."""
    out = doc_write(ws, "short.docx", title="Note", body="# Findings\n\nAll good.")
    assert out["warnings"] == []


def test_a_real_document_is_never_flagged(ws):
    out = doc_write(ws, "real.docx", title="Note",
                    body=f"# Findings\n\n{REAL_FINDING}\n\n# Recommendation\n\n- Re-measure\n- Inspect the seal")
    assert out["warnings"] == []
    assert out["content_chars"] > 150


def test_the_warning_names_the_likely_cause(ws):
    """A reader whose deck came out empty needs to know WHERE to look, and it is
    almost never the document tool - it is the step that fed it."""
    out = doc_write(ws, "hollow.pptx", title="R", body="A:\n\nB:\n\nC:")
    joined = " ".join(out["warnings"])
    assert "{{stepN}}" in joined and "FOUND" in joined


@pytest.mark.parametrize("name", ["x.docx", "x.pptx"])
def test_every_format_that_holds_prose_measures_its_own_content(ws, name):
    out = doc_write(ws, name, title="T", body="# A\n\n# B\n\n# C")
    assert out["content_chars"] == 0
    assert any("outline" in w for w in out["warnings"])


# --------------------------------------------- a passage is not a slide title

LONG_PASSAGE = (
    "During the turnaround inspection of the Crude Distillation Unit (CDU-II), "
    "comprehensive Non-Destructive Testing was performed on Crude Feed Booster "
    "Pumps P-101A/B and associated piping, and the results are summarised below."
)


def test_a_retrieved_passage_does_not_become_a_slide_title(ws):
    """"A blank-line-separated block starts a slide" was written for label lines
    like "Findings", and every existing plan writes those. Applied to a
    400-character retrieved passage it made the passage the TITLE, truncated at
    90 characters with the rest repeated underneath - so a deck built from real
    corpus text opened on slides that began mid-sentence."""
    from pptx import Presentation

    out = doc_write(ws, "deck.pptx", title="Briefing",
                    body=f"Findings:\n\n{LONG_PASSAGE}\n\n{LONG_PASSAGE}")
    titles = [
        slide.shapes.title.text
        for slide in Presentation(str(ws.resolve("deck.pptx"))).slides
        if slide.shapes.title is not None
    ]
    assert titles == ["Briefing", "Findings"], titles
    assert out["slides"] == 1, "the passages became slides of their own"


def test_a_short_label_line_still_starts_a_slide(ws):
    """The documented convention, which every existing plan depends on. A
    markup change that breaks the plans written for it is a breaking change
    wearing an improvement."""
    out = doc_write(ws, "deck.pptx", title="T",
                    body="Findings\n\n- a finding\n\nActions\n\n- an action")
    assert out["slides"] == 2


def test_no_slide_title_is_a_truncated_sentence(ws):
    from pptx import Presentation

    doc_write(ws, "deck.pptx", title="T", body=f"Findings:\n\n{LONG_PASSAGE}")
    for slide in Presentation(str(ws.resolve("deck.pptx"))).slides:
        title = slide.shapes.title
        if title is not None:
            assert not title.text.endswith("..."), f"truncated title: {title.text!r}"


# ------------------------------------------------- an outline fails the step

def test_an_outline_fails_the_step_so_the_model_rewrites_it(ws):
    """As a warning it was written, passed, and counted as a finished task. As
    a failure the model is told "this is an outline" and gets to write the
    findings - which is the only route from a hollow deck to a real one.

    The file is still written either way: refusing would lose the little
    content there is, and the reader can see it for themselves.
    """
    from kavach.agent.critic import Critic
    from kavach.types import Observation, Step, StepStatus

    out = doc_write(ws, "hollow.pptx", title="R", body="A:\n\nB:\n\nC:")
    step = Step(idx=0, intent="write the briefing", tool="doc.write", args={},
                acceptance="hollow.pptx is created in the workspace",
                status=StepStatus.RUNNING)
    verdict = Critic(semantic=False).judge(
        step,
        Observation(step_idx=0, ok=True, artifacts=["hollow.pptx"], data=out),
        ws,
    )
    assert verdict.ok is False
    assert "outline" in verdict.reason
    assert ws.resolve("hollow.pptx").exists(), "the file must still be written"


def test_a_real_document_passes_the_same_check(ws):
    from kavach.agent.critic import Critic
    from kavach.types import Observation, Step, StepStatus

    out = doc_write(ws, "real.docx", title="Note",
                    body=f"# Findings\n\n{REAL_FINDING}\n\n# Actions\n\n- Re-measure the drive end")
    step = Step(idx=0, intent="write it", tool="doc.write", args={},
                acceptance="real.docx is created in the workspace",
                status=StepStatus.RUNNING)
    verdict = Critic(semantic=False).judge(
        step, Observation(step_idx=0, ok=True, artifacts=["real.docx"], data=out), ws)
    assert verdict.ok is True
